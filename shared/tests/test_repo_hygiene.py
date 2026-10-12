"""Repository-wide checks that apply to every module."""
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import hygiene  # noqa: E402


def rel(p):
    return p.relative_to(hygiene.REPO_ROOT)


def test_no_sensitive_values_in_any_text_file():
    for path in hygiene.text_files():
        found = hygiene.findings_in_text(path.read_text(encoding="utf-8"))
        assert not found, f"{found} in {rel(path)}"


def test_no_sensitive_values_in_any_workbook_cell():
    for path in hygiene.tracked_paths():
        if path.suffix == ".xlsx":
            wb = load_workbook(path, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str):
                            found = hygiene.findings_in_text(cell.value)
                            assert not found, f"{found} in {rel(path)} {ws.title}!{cell.coordinate}"


def test_scanner_patterns_detect_what_they_claim():
    # Samples are assembled at runtime so this file never contains a literal the scanner would flag.
    samples = {
        "IPv4 address": ".".join(["10", "20", "30", "40"]),
        "Falcon CID format": "0123456789abcdef" * 2 + "-1F",
        "credential assignment": "client_" + "secret = abc123",
        "AWS access key id": "AKIA" + "ABCDEFGHIJKLMNOP",
        "private key block": "-----BEGIN RSA " + "PRIVATE KEY-----",
        "GitHub token": "ghp_" + "a" * 36,
        "JWT": "eyJhbGciOiJIUzI1NiJ9" + "." + "eyJzdWIiOiIxMjM0NTY3ODkwIn0",
    }
    assert {label for _, label in hygiene.SENSITIVE} == set(samples), "pattern without a sample"
    for pattern, label in hygiene.SENSITIVE:
        assert re.search(pattern, samples[label]), f"pattern for {label} does not match its sample"


def test_every_relative_markdown_link_resolves():
    for md in hygiene.tracked_paths():
        if md.suffix == ".md":
            bad = hygiene.broken_links(md)
            assert not bad, f"{rel(md)}: broken link(s) {bad}"


def test_gitignore_blocks_secret_files():
    ignore = (hygiene.REPO_ROOT / ".gitignore").read_text(encoding="utf-8")
    for needle in (".env", "*.pem", "*.key"):
        assert needle in ignore


def test_every_module_is_listed_in_root_readme_and_has_a_readme():
    root_readme = (hygiene.REPO_ROOT / "README.md").read_text(encoding="utf-8")
    modules = sorted(p for p in hygiene.REPO_ROOT.iterdir() if p.is_dir() and re.match(r"\d\d-", p.name))
    assert modules, "no numbered modules found"
    for m in modules:
        assert (m / "README.md").exists(), f"{m.name} has no README.md"
        assert f"({m.name})" in root_readme, f"{m.name} is not linked from the root README"
